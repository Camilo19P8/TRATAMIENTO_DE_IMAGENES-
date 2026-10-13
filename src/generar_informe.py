"""Genera el informe editable usando resultados medidos y evidencias reales."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT/'resultados/resumen.json').read_text(encoding='utf-8'))
FIG = ROOT/'resultados/figuras'
evidence=ROOT/'resultados/evidencias_vscode'
DOC = Document()
sec = DOC.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = Inches(.5)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = DOC.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0,0,0)
    for fonts in style.element.iter(qn('w:rFonts')):
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+attr),None)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_after = Pt(0)
DOC.styles['Normal'].paragraph_format.first_line_indent = Inches(.5)
DOC.styles['Title'].font.bold = True
DOC.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
DOC.styles['Heading 1'].font.bold = True
DOC.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
DOC.styles['Heading 1'].paragraph_format.space_after = Pt(6)
DOC.styles['Heading 1'].paragraph_format.space_before = Pt(0)
for style in DOC.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
header._p.append(field)
DOC.core_properties.title = 'Informe de resultados de tratamiento de imagen'
DOC.core_properties.author = 'Cristian Camilo Alarcon Delgadillo; Brayan Arley Saenz Cortes; Jonathan David Llanceros Bello'

def hyperlink(p,url):
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');props.append(fonts)
    run.append(props);txt=OxmlElement('w:t');txt.text=url;run.append(txt);link.append(run);p._p.append(link)

def para(text, indent=True, single=False):
    p=DOC.add_paragraph()
    if text.startswith('https://'):hyperlink(p,text)
    else:p.add_run(text)
    if not indent:p.paragraph_format.first_line_indent=Inches(0)
    if single:p.paragraph_format.line_spacing=1.15
    return p

def heading(text):DOC.add_heading(text,1)
def page(text):
    p=DOC.add_heading(text,1)
    p.paragraph_format.page_break_before=True
def fmt(n):return f'{n:.2f}'.replace('.',',')

def figure(num,title,path,width=6.3,note='Elaboración propia a partir de las fotografías suministradas.'):
    p=para(f'Figura {num}',False,True);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    p=para(title,False,True);p.runs[0].italic=True;p.paragraph_format.keep_with_next=True
    p=DOC.add_paragraph();p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(width))
    p=para('Nota. '+note,False,True)
    for r in p.runs:r.font.size=Pt(10)
    p.paragraph_format.space_after=Pt(8)

def table(num,title,headers,rows,widths):
    p=para(f'Tabla {num}',False,True);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    p=para(title,False,True);p.runs[0].italic=True;p.paragraph_format.keep_with_next=True
    t=DOC.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'8');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DBE7EF' if i==0 else ('F4F7F9' if i%2==0 else 'FFFFFF'));props.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.1
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(10);r.bold=(i==0)
    para('',False,True)

# Portada académica siguiendo la referencia suministrada.
p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(66)
DOC.add_paragraph('Informe de resultados de tratamiento de imagen',style='Title')
p=para('Informe de práctica · Semana 8',False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(24)
for text in ['Cristian Camilo Alarcon Delgadillo','Brayan Arley Saenz Cortes','Jonathan David Llanceros Bello']:
    p=para(text,False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=DOC.add_paragraph();p.paragraph_format.space_after=Pt(24)
for text in ['Universidad de Cundinamarca','Especialización en Inteligencia Artificial','Percepción Computacional','Mónica María Fonseca Vigoya','12 de octubre de 2026']:
    p=para(text,False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER

page('Propósito y metodología')
para('El diagnóstico de cinco fotografías de una misma escena permitió identificar una pérdida marcada de definición en la imagen I3. Aplicamos una máscara de enfoque moderada en luminancia y observamos un aumento de la respuesta de los bordes conservados. El detalle fino de las etiquetas continuó degradado; por ello, el resultado se interpreta como un realce visual parcial y no como una restauración completa.')
para('La práctica estudia cómo el encuadre, la iluminación y la estabilidad de la cámara afectan una imagen utilizada en percepción computacional. La lata negra Monster situada en la fila superior es el objeto común de las cinco tomas. Las demás latas, el teclado, las figuras y el computador aportan contexto y distractores. Los cinco archivos proceden de la carpeta de imágenes facilitada para esta actividad y se conservan sin modificar.')
para('Las condiciones se describen a partir de los nombres de los archivos y de la evidencia visible. No se dispone de registros de exposición, distancia, iluminación en lux ni dispositivo de captura. En particular, la condición denominada baja iluminación no se confirma mediante el brillo observado. Esta discrepancia se documenta para distinguir el nombre del archivo de su diagnóstico efectivo.')
para('Ejecutamos Python y OpenCV en el equipo local y organizamos el código en un proyecto de VS Code. El cuaderno complementario está preparado para ejecutar la misma práctica en Google Colab. Medimos resolución, canales, brillo, contraste, nitidez, formato y tamaño; comparamos los resultados y seleccionamos un tratamiento relacionado con el problema de I3. Las figuras muestran originales, regiones del objeto, histogramas y resultados antes y después.')

page('Fotografías de referencia e iluminación')
figure(1,'I1 con encuadre frontal y definición de referencia',FIG/'01_captura.png',5.25,'El rectángulo amarillo identifica la lata negra usada como objeto común. La escena incluye doce latas y elementos secundarios.')
figure(2,'I2 con mayor brillo observado',FIG/'02_captura.png',5.25,'Archivo original denominado baja iluminación. La imagen presenta mayor brillo que I1; el nombre no permite establecer la exposición física de la captura.')

page('Fotografías con desenfoque y fondo complejo')
figure(3,'I3 con bordes suaves y pérdida de detalle',FIG/'03_captura.png',5.25,'Se observan letras y contornos difusos. La fotografía por sí sola no permite separar con certeza desenfoque óptico y movimiento.')
figure(4,'I4 con múltiples distractores y menor escala del objeto',FIG/'04_captura.png',5.25,'El computador, la alfombrilla y los accesorios aumentan el contenido ajeno al objeto. Un contraste global alto no elimina esta dificultad.')

page('Perspectiva y comparación del mismo objeto')
figure(5,'I5 desde un ángulo elevado',FIG/'05_captura.png',5.8,'Las tapas son más visibles y cambia la proyección de la lata. La perspectiva modifica bordes, proporciones y zonas observables.')
figure(6,'Regiones del objeto en las cinco condiciones',FIG/'06_mismo_objeto.png',6.3,'De izquierda a derecha, I1 a I5. Recortes manuales normalizados a 128 × 256 píxeles con INTER_AREA. Esta normalización es exploratoria y no alinea la perspectiva.')

page('Resultados del diagnóstico')
rows=DATA['diagnostico']
table(1,'Propiedades de los archivos originales',['Imagen','Resolución\npx','Canales','Formato','Tamaño\nKiB'],[[f'I{i+1}',f"{r['ancho']} × {r['alto']}",r['canales'],r['formato'],fmt(r['bytes']/1024)] for i,r in enumerate(rows)],[.65,1.75,.85,.95,1.8])
para('Nota. KiB = bytes / 1024. Todas las imágenes se decodificaron en tres canales BGR y profundidad de 8 bits. I3 tiene una fila menos que las otras tomas; esta diferencia no explica la pérdida de definición.',False,True)
table(2,'Descriptores globales de brillo contraste y nitidez',['Imagen','Brillo\nμ','Contraste\nσ','Laplaciano\nVar','Tenengrad'],[[f'I{i+1}',fmt(r['brillo']),fmt(r['contraste']),fmt(r['laplaciano']),fmt(r['tenengrad'])] for i,r in enumerate(rows)],[.65,1.15,1.3,1.45,1.45])
para('El brillo es la media del gris OpenCV, Ygris = 0,114B + 0,587G + 0,299R, en escala de 0 a 255. El contraste es su desviación estándar. La nitidez se aproxima mediante la varianza del Laplaciano con ksize = 1 y la energía media de gradientes Sobel de tamaño 3, denominada Tenengrad. Se calculan sin redimensionar las fotografías.')
para('El ruido, la textura y la escala también afectan estos indicadores; se interpretan junto con la evidencia visual.')

page('Interpretación y selección del problema')
figure(7,'Distribución tonal y respuesta de bordes por condición',FIG/'07_diagnostico.png',5.8,'El panel derecho usa las regiones normalizadas, mientras que la tabla 2 utiliza las imágenes completas.')
table(3,'Varianza del Laplaciano en el objeto normalizado',['Imagen','I1','I2','I3','I4','I5'],[['Varianza']+[fmt(r['roi_laplaciano']) for r in rows]],[1.1,1.08,1.08,1.08,1.08,1.08])
para(f"I3 registra la menor respuesta global del Laplaciano ({fmt(rows[2]['laplaciano'])}) y también la menor respuesta en la región normalizada ({fmt(rows[2]['roi_laplaciano'])}). La coincidencia con los contornos visualmente difusos justifica seleccionarla. La reducción frente a I1 es del {fmt((1-rows[2]['laplaciano']/rows[0]['laplaciano'])*100)} % en la medida global; esta diferencia combina degradación y cambios de escena.")
para(f"I2 alcanza un brillo medio de {fmt(rows[1]['brillo'])}, frente a {fmt(rows[0]['brillo'])} en I1. El gris medio del objeto también aumenta de {fmt(rows[0]['roi_brillo'])} a {fmt(rows[1]['roi_brillo'])}. La imagen disponible no evidencia una subexposición respecto a la referencia. La autoexposición o una iluminación distinta podrían explicar el resultado, pero no se verificaron esas causas.")
para('I4 presenta distractores y menor tamaño aparente de la lata. I5 modifica la perspectiva y obtiene una alta respuesta del Laplaciano en el recorte, aunque ello no demuestra mejor legibilidad. Para esas condiciones, una segmentación o una rectificación geométrica respondería mejor al problema que aplicar un filtro de enfoque indiscriminadamente.')

page('Lectura y validación de los archivos')
figure(8,'Lectura y exportación de imágenes en VS Code',evidence/'03_lectura.jpg',6.3,'Captura real del archivo practica.py. La lectura valida que el archivo pueda decodificarse antes de medirlo.')
para('La primera etapa consiste en localizar las cinco fotografías y convertir sus archivos JPEG en matrices de píxeles. La función leer combina la lectura binaria de NumPy con la decodificación de OpenCV. Esta decisión evita dificultades con rutas que contienen tildes en Windows. Si un archivo no puede abrirse, la validación detiene el proceso y señala cuál falló.')
para('Con las matrices disponibles se obtiene el ancho, el alto y el número de canales. El formato se identifica en el archivo y el tamaño corresponde a sus bytes en disco. El procesamiento utiliza copias en memoria; las imágenes originales permanecen intactas. La función guardar codifica el resultado de salida y comprueba que la operación haya tenido éxito.')

page('Cálculo de los indicadores de diagnóstico')
figure(9,'Funciones de medición abiertas en VS Code',evidence/'04_metricas.jpg',6.3,'Captura real de los indicadores usados en el diagnóstico. Se muestran brillo, contraste, Laplaciano, Tenengrad y proporciones de intensidades extremas.')
para('La conversión a gris permite medir todas las fotografías sobre una misma escala de intensidad. La media resume el brillo y la desviación estándar describe la dispersión tonal. Los gradientes Sobel cuantifican transiciones de intensidad en dos direcciones; su energía conjunta aporta una medida complementaria a la varianza del Laplaciano.')
para('La función devuelve los valores en un registro que se integra en la tabla comparativa. También se calculan las proporciones de píxeles próximos al negro y al blanco. Estos porcentajes se conservan en los resultados digitales y ayudan a detectar extremos tonales, aunque no identifican por sí solos la causa física de una exposición o una pérdida de detalle.')

page('Separación entre medición y tratamiento')
figure(10,'Imagen original gris de diagnóstico y luminancia de tratamiento',FIG/'11_color_gris_luminancia.png',6.3,'Los paneles proceden de I3 y conservan el mismo encuadre. La imagen en gris se utiliza para medir; el canal Y permite modificar luminancia conservando los canales de crominancia durante el realce.')
para('El panel en color conserva la información visual completa de la escena. El panel gris representa la intensidad utilizada por los indicadores, mientras que el tercero muestra la luminancia Y del espacio YCrCb. En esta fotografía ambos paneles de intensidad tienen apariencia similar, pero cumplen funciones distintas en el procedimiento.')
para('Medir sobre gris facilita comparar brillo, contraste y bordes. Tratar solamente Y reduce la modificación directa de la información de color frente a realzar cada canal BGR por separado. Después del tratamiento se recombinan los canales para obtener una imagen en color con el mismo tamaño. La conversión entre espacios y la cuantización de 8 bits pueden introducir pequeñas variaciones.')
para('La comparación de los paneles también permite reconocer que la pérdida de detalle ya está presente antes del filtrado. Convertir a gris o separar luminancia no restaura los trazos difusos; estas operaciones preparan la representación que utiliza cada etapa del análisis.')

page('Construcción de la máscara de enfoque')
figure(11,'Luminancia suavizada y diferencia que refuerza los bordes',FIG/'12_etapas_luminancia.png',5.5,'Región nativa de la lata negra. El residual se representa como 128 + 6(Y − suavizada), con cero en gris medio. La ganancia de seis se usa solo para visualizar; no es la cantidad de tratamiento.')
para('El primer panel muestra la luminancia original del objeto. El segundo aplica un suavizado gaussiano con sigma de 1,2 píxeles: se atenúan pequeñas variaciones y transiciones. Restar esa versión suavizada de la luminancia original produce una diferencia que concentra parte de la información de bordes y textura.')
para('El tercer panel representa esa diferencia mediante una escala centrada en gris. Los valores positivos se ven más claros y los negativos, más oscuros. El aumento visual permite localizar trazos que serían difíciles de distinguir en el residual sin amplificación. La imagen sirve para explicar la operación y no constituye una restauración adicional.')

page('Tratamiento aplicado en VS Code')
para('Seleccionamos una máscara de enfoque sobre el canal Y del espacio YCrCb. El desenfoque atenúa transiciones entre intensidades; sumar una fracción de la diferencia entre la luminancia original y su versión suavizada refuerza los bordes que todavía permanecen. Usar únicamente Y permite conservar los canales de crominancia durante la operación. La conversión y la cuantización pueden producir pequeñas diferencias de color.')
para('La expresión utilizada es Y′ = clip[Y + a(Y − GaussianBlur(Y, σ)), 0, 255], con cantidad a = 0,8 y desviación σ = 1,2 píxeles. El tamaño del núcleo se calcula automáticamente mediante (0, 0). El recorte se mide antes de limitar los valores y se exporta el resultado en PNG para evitar una nueva compresión JPEG con pérdida.')
figure(12,'Implementación del realce sobre luminancia en VS Code',evidence/'05_realce.jpg',6.3,'Captura real de la función realzar. Se muestran la conversión, el suavizado, el cálculo de la diferencia, la limitación del rango y la reconstrucción del color.')

page('Exploración y selección de los parámetros')
figure(13,'Evaluación de tres cantidades de realce en VS Code',evidence/'06_parametros.jpg',6.3,'El mismo procedimiento evalúa cantidades de 0,4, 0,8 y 1,2 con sigma fijo. Cada opción se mide y se exporta antes de elegir la cantidad intermedia.')
table(4,'Exploración de la cantidad de realce con sigma fijo',['Cantidad a','Laplaciano\nglobal','Tenengrad\nglobal','Recorte Y\n%'],[[fmt(r['cantidad']),fmt(r['laplaciano']),fmt(r['tenengrad']),f"{r['recorte_previo_pct']:.4f}".replace('.',',')] for r in DATA['parametros']],[1.1,1.8,1.8,1.3])
para('Elegimos 0,8 como realce moderado. La cantidad 1,2 casi duplica el recorte; la comparación visual siguiente permite valorar su efecto.')

page('Comparación visual de las cantidades de realce')
figure(14,'Mismo objeto con tres intensidades del tratamiento',FIG/'14_parametros_detalle.png',6.3,'Todos los paneles corresponden al mismo recorte nativo de I3. El panel original permite distinguir el efecto del filtro de un cambio de encuadre.')
para('La cantidad 0,4 introduce un realce suave. La opción 0,8 acentúa más los contornos del logotipo y del cuerpo de la lata, y la cantidad 1,2 intensifica aún más las transiciones. La lectura visual debe realizarse sobre las mismas letras y bordes: un panel más contrastado no significa que contenga información nueva.')
para('La tabla 4 complementa esta comparación al mostrar el crecimiento de la respuesta a bordes y del recorte de luminancia. Elegimos la cantidad intermedia para moderar el efecto. La comparación no demuestra un máximo de calidad ni un parámetro óptimo para cualquier fotografía; permite justificar una decisión específica para la imagen analizada.')
para('El texto fino continúa degradado en las tres alternativas. Aumentar la cantidad puede reforzar artefactos o transiciones no deseadas. La ecualización y CLAHE se descartaron porque responden a un problema tonal diferente (OpenCV, s. f.-a). La selección conserva como conclusión un realce parcial de los bordes presentes.')

page('Comparación visual antes y después')
figure(15,'Resultado del tratamiento en la imagen completa',FIG/'08_antes_despues.png',6.3,'Ambos paneles corresponden a I3; se conserva la misma resolución y el mismo encuadre.')
figure(16,'Detalle de la lata negra en la misma región de I3',FIG/'09_detalle.png',4.4,'Recorte nativo de 125 × 240 píxeles en ambos paneles, ampliado únicamente para visualizarlo. El tratamiento se realizó sobre la imagen completa.')
para('El borde de la lata y algunos trazos del logotipo se perciben más definidos. Las letras pequeñas siguen borrosas y algunas transiciones se acentúan. Se observa un realce moderado, sin evidencia de recuperación de información perdida.')

page('Inspección del resultado en dos detalles de la escena')
figure(17,'Comparación de texto y contornos en la lata y el sobre',FIG/'15_detalles_resultado.png',5.4,'Cada fila presenta una región nativa de I3 antes y después. Los recortes solo se amplían para observarlos; la intervención se aplica a la fotografía completa.')
para('La primera fila permite revisar los trazos del logotipo de la lata, que adquieren transiciones más marcadas. La segunda muestra el sobre situado frente al teclado, donde las letras y bordes siguen afectados por el desenfoque. La comparación amplía la evaluación a otro elemento de la misma escena sin cambiar el tratamiento.')
para('El resultado es coherente en ambos detalles: se refuerzan contornos conservados, mientras que la información fina permanece incompleta. Esta evidencia apoya la recomendación de repetir la captura cuando se necesite leer texto con precisión. No se atribuye al filtro una mejora de OCR, porque esa tarea no se ejecutó.')

page('Respuesta de bordes y alcance del cambio')
figure(18,'Magnitud del Laplaciano antes y después en el objeto',FIG/'13_respuesta_bordes.png',5.4,'Se utiliza la misma escala de visualización en ambos paneles, desde cero hasta el máximo conjunto de las dos matrices. Los valores absolutos muestran magnitud, sin conservar el signo del Laplaciano.')
para('El Laplaciano responde a cambios de intensidad en la región analizada. Tras el realce, algunas transiciones adquieren mayor magnitud y aparecen más claras en el panel derecho. Esta representación ayuda a interpretar por qué aumenta la varianza registrada en las tablas.')
para('La escala común evita que cada panel se normalice independientemente y produzca una apariencia engañosa de similitud. Aun así, una respuesta más intensa no demuestra una restauración más fiel: el filtro modifica justamente esas transiciones. La evaluación de detalle y de legibilidad sigue siendo necesaria para interpretar el cambio.')

page('Evaluación numérica del resultado')
before,after=DATA['antes'],DATA['despues']
table(5,'Comparación global de I3 antes y después',['Indicador','Antes','Después','Cambio\n%'],[[label,fmt(before[key]),fmt(after[key]),fmt((after[key]/before[key]-1)*100)] for label,key in [('Brillo medio','brillo'),('Contraste','contraste'),('Varianza Laplaciano','laplaciano'),('Tenengrad','tenengrad')]], [2.3,1.3,1.3,1.1])
figure(19,'Efecto sobre el histograma y la respuesta a bordes',FIG/'10_evaluacion.png',6.3,'La distribución tonal cambia poco, mientras que el realce aumenta la respuesta del Laplaciano.')
para(f"La varianza del Laplaciano aumenta de {fmt(before['laplaciano'])} a {fmt(after['laplaciano'])}, un {fmt((after['laplaciano']/before['laplaciano']-1)*100)} %. Tenengrad aumenta un {fmt((after['tenengrad']/before['tenengrad']-1)*100)} %. El brillo disminuye solo {fmt(before['brillo']-after['brillo'])} niveles y el contraste cambia un {fmt((after['contraste']/before['contraste']-1)*100)} %. El comportamiento concuerda con un realce de bordes y un ajuste tonal pequeño.")
recorte_texto=f"{DATA['recorte_previo_pct']:.4f}".replace('.',',')
para(f"En el recorte nativo de la lata, el Laplaciano pasa de {fmt(DATA['antes_roi']['laplaciano'])} a {fmt(DATA['despues_roi']['laplaciano'])}. Estos valores no se comparan directamente con la tabla 3, que utiliza recortes normalizados. El {recorte_texto} % de valores Y se recorta antes de reconstruir el color; esta medida es distinta de la proporción de píxeles claros o negros del gris final.")

page('Cambios de intensidad en la región del objeto')
figure(20,'Histogramas y diferencias de intensidad en el recorte nativo',FIG/'16_cambios_region.png',6.3,'El panel derecho calcula la diferencia del gris final entre los mismos píxeles después y antes. No se compara con la fotografía de referencia ni con otra perspectiva.')
para('El histograma del panel izquierdo resume cómo se distribuyen las intensidades dentro de la lata seleccionada. El segundo panel muestra cuántos píxeles aumentan o disminuyen su intensidad. Las diferencias positivas representan aclaramiento y las negativas, oscurecimiento en esas posiciones.')
para(f"En esta región, el brillo medio pasa de {fmt(DATA['antes_roi']['brillo'])} a {fmt(DATA['despues_roi']['brillo'])}, mientras que el contraste pasa de {fmt(DATA['antes_roi']['contraste'])} a {fmt(DATA['despues_roi']['contraste'])}. El efecto sobre la media es pequeño y la dispersión aumenta. Esta combinación concuerda con un realce que modifica transiciones y no intenta cambiar globalmente la iluminación.")
para('La coincidencia de encuadre y tamaño permite asociar la diferencia por píxel al procedimiento aplicado. La gráfica no expresa error de restauración, porque la imagen antes del filtro ya contiene degradación. Para medir fidelidad se necesitaría una referencia nítida alineada de ese mismo encuadre.')

page('Limitaciones conclusiones y recomendaciones')
para('La máscara de enfoque mejora parcialmente la apariencia de los bordes presentes. La inspección de las figuras 16 y 17 muestra que el texto fino continúa degradado, por lo que el tratamiento no satisface por sí solo una tarea de lectura precisa de etiquetas. Las medidas de bordes aumentan en parte por la propia definición del filtro; también podrían aumentar por ruido o halos. No se evaluó un clasificador, un detector ni un sistema OCR, y no se afirma una mejora de exactitud en esas tareas.')
para('Las cinco tomas presentan cambios de escala, distancia aparente, iluminación y perspectiva a la vez. El recorte del objeto reduce la influencia del fondo, pero la normalización modifica el muestreo y no proporciona correspondencia punto a punto. No se calcularon PSNR ni SSIM entre I1 e I3, porque no existe una imagen nítida alineada del mismo encuadre que permita atribuir esas métricas a restauración.')
para('La fotografía identificada como baja iluminación requiere una nueva captura controlada para demostrar esa condición. Se recomienda mantener cámara, distancia y encuadre fijos, reducir la luz y registrar la exposición. Suprimir el autoajuste de exposición cuando el equipo lo permita facilitaría relacionar el cambio de luz con el resultado. Una imagen artificialmente oscurecida serviría como simulación, pero no sustituiría la fotografía propia solicitada.')
para('Para el desenfoque, la mejor medida es repetir la captura apoyando la cámara y enfocando sobre la lata. Para I4 conviene reducir distractores o segmentar el objeto; para I5, documentar la perspectiva y, si la tarea lo exige, estimar una transformación geométrica. La práctica evidencia que cada intervención debe partir del problema observado y valorarse con imágenes, indicadores y limitaciones explícitas.')

page('Exportación y trazabilidad de los resultados')
figure(21,'Exportación de métricas y huellas de los originales en VS Code',evidence/'07_exportacion.jpg',6.3,'Captura real del código que conserva resultados numéricos, versiones del entorno y un manifiesto de las fotografías.')
para('El proceso guarda la comparación antes y después y los parámetros evaluados en archivos CSV. El resumen JSON reúne los valores y las versiones del entorno para revisarlos con precisión. El manifiesto registra una huella SHA-256 de cada original y las coordenadas de la región seleccionada, lo que permite comprobar que las imágenes de entrada no cambiaron.')
para('Las figuras del proceso se conservan junto con las capturas reales del editor. El código y estas evidencias permiten revisar cada operación y repetirla con las mismas entradas. El informe Word y su PDF se entregan como documentos separados del repositorio de la práctica.')

page('Resultados revisables y reproducibilidad')
if (evidence/'08_resultados_ampliados.jpg').exists():
    figure(22,'Detalles del resultado abiertos en VS Code',evidence/'08_resultados_ampliados.jpg',6.5,'Captura real del visor de imágenes de VS Code. Los paneles proceden de las regiones de la lata y el sobre medidas en la misma fotografía tratada.')
para(f"La ejecución local utilizó Python {DATA['python']}, OpenCV {DATA['opencv']}, NumPy {DATA['numpy']} y Matplotlib {DATA['matplotlib']}. El procesamiento conserva los originales y utiliza parámetros fijos. El archivo manifiesto.json registra las huellas SHA-256 y las coordenadas de cada región del objeto.")
para('Repositorio y código de la práctica',False,True)
para('https://github.com/Camilo19P8/TRATAMIENTO_DE_IMAGENES-',False,True)
para('Cuaderno preparado para Google Colab',False,True)
para('https://colab.research.google.com/github/Camilo19P8/TRATAMIENTO_DE_IMAGENES-/blob/main/notebooks/Practica_tratamiento_imagenes.ipynb',False,True)
para('Para reproducir, ejecutar practica.py y luego proceso_visual.py con las dependencias de requirements.txt. El README detalla el procedimiento. La ejecución documentada es local; el cuaderno permite repetir el análisis en Colab. El informe se entrega separado del repositorio.',False,True)

page('Referencias y correspondencia con la actividad')
for author,title,container,url in [
('American Psychological Association. (s. f.). ','Sample student paper annotated','', 'https://apastyle.apa.org/style-grammar-guidelines/paper-format/student-annotated.pdf'),
('OpenCV. (s. f.-a). ','Histograms - 2: Histogram equalization','OpenCV 4.13.0 documentation. ', 'https://docs.opencv.org/4.13.0/d5/daf/tutorial_py_histogram_equalization.html'),
('OpenCV. (s. f.-b). ','Smoothing images','OpenCV 4.13.0 documentation. ', 'https://docs.opencv.org/4.13.0/d4/d13/tutorial_py_filtering.html')]:
    p=para(author,False);p.add_run(title).italic=True;p.add_run('. '+container);hyperlink(p,url)
    p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
para('El filtrado gaussiano utilizado se apoya en las funciones descritas por OpenCV (s. f.-b). El formato académico sigue el modelo de trabajo estudiantil de APA. Se conservan divisiones internas visibles en las tablas por solicitud expresa de presentación; esta adaptación se aparta de la convención APA de utilizar pocas líneas horizontales y evitar líneas verticales.',False)
table(6,'Localización de los elementos solicitados',['Elemento','Evidencia en el informe'],[
['Cinco fotografías propias','Figuras 1 a 5; originales en datos/originales'],
['Diagnóstico y tabla comparativa','Tablas 1 a 3; figura 7; diagnostico.csv'],
['Problema y tratamiento justificado','Selección de I3; tabla 4; función realzar'],
['Comparación visual y numérica','Figuras 14 a 20; tabla 5; antes_despues.csv'],
['Código y enlace funcional','Figuras 8, 9, 12, 13 y 21; repositorio'],
['Análisis conclusiones y límites','Interpretación; evaluación; recomendaciones'],
['PDF de al menos dos páginas','Informe completo con portada y referencias'],
['Condición de baja iluminación','Discrepancia documentada; requiere recaptura controlada']
],[2.3,3.7])
(ROOT/'informe').mkdir(exist_ok=True)
DOC.save(ROOT/'informe/Informe_tratamiento_imagenes.docx')
print('Informe DOCX generado')
